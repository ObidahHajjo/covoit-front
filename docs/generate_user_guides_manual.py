from __future__ import annotations

import json
from pathlib import Path

from docx import Document

import generate_user_guides as base


def tr_factory(translations: dict[str, str]):
    return base.tr_factory(translations)


def chapter_intro_fr(tr):
    return base.Chapter(
        title="Introduction",
        overview="Ce manuel accompagne les utilisatrices et utilisateurs de Covoit dans la prise en main de l'application de covoiturage. Il s'adresse aux debutants comme aux personnes deja a l'aise avec les applications mobiles et web. Avant de commencer, preparez une connexion internet stable et un compte actif. Cette version du document reprend les ecrans disponibles dans l'application actuelle et pourra etre mise a jour si de nouvelles fonctions sont ajoutees.",
        steps=[
            base.StepGroup(
                heading="Identifier l'objectif de l'application",
                instructions=[
                    "Utilisez Covoit pour rechercher un trajet, proposer un trajet, reserver une place et coordonner les details pratiques dans la messagerie.",
                    "Reperez les espaces principaux du produit : connexion, tableau de bord, recherche, reservations, trajets conducteur, messages et compte.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-home-dashboard.png",
                placeholder=None,
                caption="Figure 1 - Vue generale de l'application apres connexion",
            ),
        ],
        notes="Note - Ce manuel correspond a la version actuelle de l'application visible dans l'environnement de demonstration local.",
    )


def chapter_safety_fr(tr):
    return base.Chapter(
        title="Securite et conformite",
        overview="Covoit est une application logicielle : il n'y a pas de consignes materielles de manipulation, mais certaines bonnes pratiques numeriques restent essentielles. Protegez vos identifiants, n'utilisez pas un appareil partage sans deconnexion, et verifiez les informations du trajet avant de confirmer une reservation. Les donnees personnelles visibles dans l'application doivent etre traitees avec attention conformement aux regles de votre organisation et a la legislation locale en vigueur.",
        steps=[
            base.StepGroup(
                heading="Verifier les informations sensibles",
                instructions=[
                    "Controlez votre e-mail, votre numero de telephone et les informations du conducteur avant de valider un trajet.",
                    "Ne partagez jamais votre mot de passe dans les messages ou avec une autre personne.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-account-profile.png",
                placeholder=None,
                caption="Figure 2 - Informations personnelles a verifier dans le compte",
            ),
        ],
        notes="Astuce - Utilisez le module de changement de mot de passe dans le compte pour renforcer regulierement la securite de votre acces.",
    )


def chapter_getting_started_fr(tr):
    return base.Chapter(
        title="Prise en main",
        overview="La prise en main commence par l'ouverture de l'application, la creation d'un compte si necessaire, puis la connexion a votre espace personnel. Les ecrans d'acces vous permettent aussi de recuperer votre mot de passe en cas d'oubli.",
        steps=[
            base.StepGroup(
                heading="Se connecter",
                instructions=[
                    f"Ouvrez la page {tr('auth.signIn')}.",
                    f"Saisissez {tr('auth.emailAddress')} et {tr('auth.password')}, puis cliquez sur {tr('auth.signIn')}.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-authentication-1.png",
                placeholder=None,
                caption="Figure 3 - Ecran de connexion",
            ),
            base.StepGroup(
                heading="Creer un compte",
                instructions=[
                    f"Cliquez sur {tr('auth.createAccount')} depuis l'ecran de connexion.",
                    f"Renseignez {tr('common.email')}, {tr('common.password')} et {tr('common.confirm')} avant de valider.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-authentication-2.png",
                placeholder=None,
                caption="Figure 4 - Ecran de creation de compte",
            ),
            base.StepGroup(
                heading="Recuperer l'acces",
                instructions=[
                    f"Utilisez {tr('auth.forgotPassword')} si vous ne vous souvenez plus de votre mot de passe.",
                    f"Envoyez le lien de reinitialisation puis definissez un nouveau mot de passe via {tr('auth.resetPassword')}.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-authentication-3.png",
                placeholder=None,
                caption="Figure 5 - Ecran de recuperation de mot de passe",
            ),
        ],
        notes="Note - Si vous ouvrez le lien recu par e-mail, l'application affiche ensuite un ecran de reinitialisation du mot de passe avant de vous ramener vers la connexion.",
    )


def feature_chapters_fr(tr):
    return [
        base.Chapter(
            title="Fonctionnalites - Tableau de bord",
            overview="Le tableau de bord centralise les informations utiles au quotidien : acces rapide aux trajets, reservations a venir et liens vers les principales zones de l'application.",
            steps=[
                base.StepGroup(
                    heading="Consulter la page d'accueil",
                    instructions=[
                        f"Ouvrez {tr('nav.home')} apres la connexion.",
                        f"Reperez les blocs {tr('home.driverTrips')} et {tr('home.bookings')} pour suivre votre activite.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-home-dashboard.png",
                    placeholder=None,
                    caption="Figure 6 - Tableau de bord",
                ),
            ],
            notes="Astuce - Utilisez cette page comme point d'entree principal pour gagner du temps dans vos actions frequentes.",
        ),
        base.Chapter(
            title="Fonctionnalites - Rechercher et reserver un trajet",
            overview="La recherche de trajet permet de comparer les propositions disponibles puis de consulter les details d'un trajet avant de reserver une place ou de contacter le conducteur.",
            steps=[
                base.StepGroup(
                    heading="Lancer une recherche",
                    instructions=[
                        f"Ouvrez {tr('shell.findTrips')}.",
                        f"Saisissez {tr('search.departureCity')} et {tr('search.arrivalCity')}, puis lancez la recherche.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-find-trip-form.png",
                    placeholder=None,
                    caption="Figure 7 - Formulaire de recherche",
                ),
                base.StepGroup(
                    heading="Consulter les resultats et les details",
                    instructions=[
                        "Ouvrez la page de resultats puis choisissez un trajet a examiner plus en detail.",
                        f"Depuis la fiche du trajet, utilisez {tr('trip.confirmBooking')} ou {tr('trip.contactDriver')} selon votre besoin.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-trip-details.png",
                    placeholder=None,
                    caption="Figure 8 - Details d'un trajet a reserver",
                ),
            ],
            notes="Note - Selon les donnees disponibles, la recherche peut afficher peu ou pas de resultats pour certaines villes ou certaines dates.",
        ),
        base.Chapter(
            title="Fonctionnalites - Mes reservations",
            overview="La zone des reservations rassemble vos places reservees, leur statut, ainsi que les actions utiles comme l'ouverture du detail ou la prise de contact avec le conducteur.",
            steps=[
                base.StepGroup(
                    heading="Afficher les reservations",
                    instructions=[
                        f"Ouvrez {tr('shell.bookings')} pour afficher la liste de vos trajets reserves.",
                        "Selectionnez une reservation pour verifier son statut et ses informations pratiques.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-bookings-list.png",
                    placeholder=None,
                    caption="Figure 9 - Liste des reservations",
                ),
                base.StepGroup(
                    heading="Verifier le detail d'une reservation",
                    instructions=[
                        f"Relisez {tr('bookings.reservationStatus')} et les informations du trajet.",
                        f"Utilisez {tr('trip.contactDriver')} si vous devez confirmer un detail avec le conducteur.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-booking-details.png",
                    placeholder=None,
                    caption="Figure 10 - Detail d'une reservation",
                ),
            ],
            notes="Astuce - Consultez le detail avant le depart pour confirmer l'adresse, l'horaire et l'etat du trajet.",
        ),
        base.Chapter(
            title="Fonctionnalites - Proposer et gerer mes trajets",
            overview="Les conducteurs peuvent publier un trajet, definir les informations du vehicule, choisir les adresses et gerer les passagers depuis leur espace dedie.",
            steps=[
                base.StepGroup(
                    heading="Afficher mes trajets",
                    instructions=[
                        f"Ouvrez {tr('shell.myTrips')} pour consulter vos trajets en cours, a venir et passes.",
                        f"Utilisez {tr('driverTrips.publishNew')} pour creer un nouveau trajet.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-driver-trips-list.png",
                    placeholder=None,
                    caption="Figure 11 - Liste des trajets conducteur",
                ),
                base.StepGroup(
                    heading="Publier un nouveau trajet",
                    instructions=[
                        "Renseignez la date, l'heure, les places et les adresses de depart et d'arrivee.",
                        "Choisissez les suggestions d'adresses avant de valider la publication.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-driver-publish-trip.png",
                    placeholder=None,
                    caption="Figure 12 - Formulaire de publication d'un trajet",
                ),
                base.StepGroup(
                    heading="Suivre un trajet publie",
                    instructions=[
                        "Ouvrez un trajet existant pour consulter les passagers et les actions disponibles.",
                        "Utilisez la fiche detaillee pour coordonner le depart ou annuler si necessaire.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-driver-trip-details.png",
                    placeholder=None,
                    caption="Figure 13 - Gestion d'un trajet conducteur",
                ),
            ],
            notes="Note - La publication depend de suggestions d'adresses valides et des permissions disponibles sur le compte conducteur.",
        ),
        base.Chapter(
            title="Fonctionnalites - Messages et coordination",
            overview="La messagerie integree facilite les echanges entre conducteurs et passagers pour confirmer les details du trajet sans sortir de l'application.",
            steps=[
                base.StepGroup(
                    heading="Ouvrir la boite de reception",
                    instructions=[
                        f"Ouvrez {tr('shell.chats')} pour voir les conversations disponibles.",
                        "Reperez les nouveaux messages puis ouvrez la conversation qui vous interesse.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-chat-inbox.png",
                    placeholder=None,
                    caption="Figure 14 - Boite de reception",
                ),
                base.StepGroup(
                    heading="Envoyer et gerer des messages",
                    instructions=[
                        f"Saisissez votre texte puis utilisez {tr('common.send')}.",
                        "Vous pouvez aussi selectionner, copier ou effacer des messages selon les options visibles dans la conversation.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-chat-conversation.png",
                    placeholder=None,
                    caption="Figure 15 - Conversation ouverte",
                ),
            ],
            notes="Astuce - Servez-vous de la conversation pour confirmer le lieu exact et l'heure de rendez-vous juste avant le depart.",
        ),
        base.Chapter(
            title="Fonctionnalites - Mon compte et mon vehicule",
            overview="Le compte personnel permet de mettre a jour vos informations, renforcer la securite de l'acces et renseigner le vehicule utilise pour les trajets conducteur.",
            steps=[
                base.StepGroup(
                    heading="Mettre a jour le profil",
                    instructions=[
                        f"Ouvrez {tr('shell.account')} puis l'onglet {tr('common.profile')}.",
                        "Modifiez vos informations visibles et enregistrez-les.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-account-profile.png",
                    placeholder=None,
                    caption="Figure 16 - Onglet profil",
                ),
                base.StepGroup(
                    heading="Mettre a jour le vehicule",
                    instructions=[
                        f"Passez dans l'onglet {tr('common.car')}.",
                        "Renseignez ou modifiez les informations du vehicule, puis enregistrez-les.",
                    ],
                    screenshot_path=base.SCREENSHOTS_DIR / "screenshot-account-car.png",
                    placeholder=None,
                    caption="Figure 17 - Onglet voiture",
                ),
            ],
            notes="Note - Cette zone contient aussi les actions sensibles, comme le changement de mot de passe et la suppression du compte.",
        ),
    ]


def chapter_troubleshooting_fr():
    return base.Chapter(
        title="Depannage",
        overview="Cette section aide a resoudre les problemes courants sans avoir a contacter le support. Elle couvre les situations les plus frequentes observees lors de la connexion, de la recherche de trajets et de l'affichage des donnees.",
        steps=[
            base.StepGroup(
                heading="Verifier les problemes courants",
                instructions=[
                    "Si la page ne charge pas correctement, rafraichissez l'onglet ou reconnectez-vous.",
                    "Si une liste reste vide, verifiez qu'un trajet, une reservation ou une conversation existe bien pour votre compte.",
                    "Si les suggestions d'adresses ne s'affichent pas, verifiez la connexion reseau et relancez la recherche.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-find-trip-results-empty.png",
                placeholder=None,
                caption="Figure 18 - Exemple d'ecran avec peu ou pas de resultats",
            ),
        ],
        notes="Astuce - La plupart des problemes d'affichage proviennent de donnees absentes, d'une session expiree ou d'une connexion reseau instable.",
    )


def chapter_faq_fr():
    return base.Chapter(
        title="FAQ",
        overview="Cette foire aux questions reprend les demandes les plus frequentes afin de repondre rapidement aux besoins les plus courants des utilisatrices et utilisateurs.",
        steps=[
            base.StepGroup(
                heading="Questions frequentes",
                instructions=[
                    "Pourquoi je ne vois pas certains ecrans ? Certains espaces dependent des permissions du compte ou des donnees disponibles.",
                    "Pourquoi le bouton reste desactive ? Des champs obligatoires peuvent etre vides ou invalides.",
                    "Puis-je modifier mes informations personnelles ? Oui, depuis la page Mon compte.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-account-profile.png",
                placeholder=None,
                caption="Figure 19 - Le compte aide a gerer les informations personnelles",
            ),
        ],
        notes="Note - Pour des situations plus specifiques, utilisez la section Contact et assistance a la fin du manuel.",
    )


def chapter_maintenance_fr():
    return base.Chapter(
        title="Maintenance et bonnes pratiques",
        overview="Comme il s'agit d'une application logicielle, la maintenance consiste surtout a garder vos informations a jour, a verifier votre mot de passe et a maintenir un historique clair dans vos trajets et messages.",
        steps=[
            base.StepGroup(
                heading="Entretenir son compte",
                instructions=[
                    "Mettez a jour vos informations de profil si votre nom d'usage, votre pseudo ou votre numero changent.",
                    "Verifiez aussi les informations du vehicule avant de proposer un nouveau trajet.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-account-profile.png",
                placeholder=None,
                caption="Figure 20 - Espace de maintenance du compte",
            ),
        ],
        notes="Astuce - Un compte bien renseigne rassure les autres membres et facilite la coordination avant le depart.",
    )


def chapter_support_fr():
    return base.Chapter(
        title="Contact et assistance",
        overview="Si vous avez encore besoin d'aide, preparez quelques informations utiles avant de demander une assistance technique ou organisationnelle. Cela permet un traitement plus rapide et plus precis de votre demande.",
        steps=[
            base.StepGroup(
                heading="Preparer sa demande d'assistance",
                instructions=[
                    "Notez la page concernee, l'heure du probleme et l'action que vous tentiez d'effectuer.",
                    "Conservez si possible une capture d'ecran du probleme ou du message affiche.",
                    "Transmettez ces informations au support ou a l'organisation qui vous a fourni l'acces a l'application.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-home-dashboard.png",
                placeholder=None,
                caption="Figure 21 - Exemple de page a citer dans une demande de support",
            ),
        ],
        notes="Note - Si votre deployment Covoit ne montre pas de contact direct dans l'application, rapprochez-vous de votre structure ou de l'equipe projet.",
    )


def chapter_appendix_fr():
    return base.Chapter(
        title="Annexe et glossaire",
        overview="L'annexe rappelle quelques termes utiles utilises dans l'application afin de rendre la lecture plus simple pour tous les profils d'utilisateurs.",
        steps=[
            base.StepGroup(
                heading="Glossaire rapide",
                instructions=[
                    "Trajet : deplacement propose par un conducteur.",
                    "Reservation : place reservee par un passager sur un trajet.",
                    "Conversation : fil de messages entre conducteur et passager.",
                    "Mon compte : espace de gestion des informations personnelles et du vehicule.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-home-dashboard.png",
                placeholder=None,
                caption="Figure 22 - Les espaces principaux de l'application servent de reperes dans le glossaire",
            ),
        ],
        notes="Astuce - Relisez ce glossaire si certains termes de l'interface vous semblent ambigus lors des premieres utilisations.",
    )


def chapter_index_fr():
    return base.Chapter(
        title="Index",
        overview="Cet index thematique aide a retrouver rapidement les sujets principaux du manuel, en particulier dans une version imprimee ou partagee hors de l'application Word.",
        steps=[
            base.StepGroup(
                heading="Reperes rapides",
                instructions=[
                    "Compte : voir Introduction, Maintenance et bonnes pratiques, Contact et assistance.",
                    "Connexion : voir Prise en main.",
                    "Messages : voir Fonctionnalites - Messages et coordination.",
                    "Reservations : voir Fonctionnalites - Mes reservations.",
                    "Trajets conducteur : voir Fonctionnalites - Proposer et gerer mes trajets.",
                ],
                screenshot_path=base.SCREENSHOTS_DIR / "screenshot-home-dashboard.png",
                placeholder=None,
                caption="Figure 23 - Ecran de reference pour les sujets principaux du manuel",
            ),
        ],
        notes="Note - Le sommaire automatique de Word reste le moyen le plus rapide pour naviguer vers chaque section du document.",
    )


def build_manual_fr(tr):
    return [
        chapter_intro_fr(tr),
        chapter_safety_fr(tr),
        chapter_getting_started_fr(tr),
        *feature_chapters_fr(tr),
        chapter_troubleshooting_fr(),
        chapter_faq_fr(),
        chapter_maintenance_fr(),
        chapter_support_fr(),
        chapter_appendix_fr(),
        chapter_index_fr(),
    ]


def translate_chapter(chapter: base.Chapter, translator) -> base.Chapter:
    return chapter


def build_manual_en(tr):
    return [
        base.Chapter(
            title="Introduction",
            overview="This manual helps users understand and operate the Covoit carpooling application. It is written for both first-time and regular users. Before you begin, make sure you have a stable internet connection and an active account. This version reflects the current application build and may be updated when new features are added.",
            steps=[base.StepGroup("Understand the product purpose", ["Use Covoit to find rides, offer rides, reserve seats, and coordinate details through in-app chat.", "Identify the main product areas: access, dashboard, search, bookings, driver trips, messages, and account."], base.SCREENSHOTS_DIR / "screenshot-home-dashboard.png", None, "Figure 1 - General application view after sign-in")],
            notes="Note - This manual reflects the current app version available in the local demonstration environment.",
        ),
        base.Chapter(
            title="Safety and compliance information",
            overview="Covoit is a software product, so there are no physical handling instructions, but basic digital safety remains important. Protect your credentials, avoid leaving your session open on a shared device, and review ride information carefully before confirming a booking. Personal information shown in the app should be handled according to your organization rules and local privacy requirements.",
            steps=[base.StepGroup("Review sensitive information", ["Check your email, phone number, and driver details before confirming a ride.", "Never share your password in chat or with another user."], base.SCREENSHOTS_DIR / "screenshot-account-profile.png", None, "Figure 2 - Personal information to review in the account area")],
            notes="Tip - Use the password change module in the account area to strengthen account access regularly.",
        ),
        base.Chapter(
            title="Getting started",
            overview="Getting started begins with opening the application, creating an account when needed, and signing in to your personal space. The access screens also let you recover your password if you forget it.",
            steps=[
                base.StepGroup("Sign in", [f"Open the {tr('auth.signIn')} page.", f"Enter {tr('auth.emailAddress')} and {tr('auth.password')}, then click {tr('auth.signIn')}."], base.SCREENSHOTS_DIR / "screenshot-authentication-1-en.png", None, "Figure 3 - Sign-in screen"),
                base.StepGroup("Create an account", [f"Click {tr('auth.createAccount')} from the sign-in screen.", f"Fill in {tr('common.email')}, {tr('common.password')}, and {tr('common.confirm')} before you continue."], base.SCREENSHOTS_DIR / "screenshot-authentication-2-en.png", None, "Figure 4 - Create account screen"),
                base.StepGroup("Recover access", [f"Use {tr('auth.forgotPassword')} if you no longer know your password.", f"Send the reset link and complete the new password screen through {tr('auth.resetPassword')}."], base.SCREENSHOTS_DIR / "screenshot-authentication-3-en.png", None, "Figure 5 - Password recovery screen"),
            ],
            notes="Note - When you open the email link, the app shows a password reset screen before returning you to sign-in.",
        ),
        base.Chapter("Features and functions - Home dashboard", "The dashboard centralizes everyday information: quick access to rides, upcoming bookings, and direct links to the main app areas.", [base.StepGroup("Open the home page", [f"Open {tr('nav.home')} after sign-in.", f"Look for {tr('home.driverTrips')} and {tr('home.bookings')} blocks to track your activity."], base.SCREENSHOTS_DIR / "screenshot-home-dashboard.png", None, "Figure 6 - Home dashboard")], "Tip - Use this page as the main starting point for frequent actions."),
        base.Chapter("Features and functions - Find and book a ride", "Ride search helps users compare available options, review ride details, and then reserve a seat or contact the driver.", [base.StepGroup("Start a search", [f"Open {tr('shell.findTrips')}.", f"Enter {tr('search.departureCity')} and {tr('search.arrivalCity')}, then run the search."], base.SCREENSHOTS_DIR / "screenshot-find-trip-form.png", None, "Figure 7 - Search form"), base.StepGroup("Review results and ride details", ["Open the results page and choose one ride to inspect.", f"From the ride page, use {tr('trip.confirmBooking')} or {tr('trip.contactDriver')} depending on your need."], base.SCREENSHOTS_DIR / "screenshot-trip-details.png", None, "Figure 8 - Ride details before booking")], "Note - Depending on available data, some searches may show only a few rides or none at all."),
        base.Chapter("Features and functions - My bookings", "The bookings area gathers reserved seats, their status, and the actions available to the passenger.", [base.StepGroup("Open the bookings list", [f"Open {tr('shell.bookings')} to display your reserved rides.", "Select one booking to check its status and practical details."], base.SCREENSHOTS_DIR / "screenshot-bookings-list.png", None, "Figure 9 - Bookings list"), base.StepGroup("Review one booking", [f"Read {tr('bookings.reservationStatus')} and review the ride information.", f"Use {tr('trip.contactDriver')} when you need to confirm a detail with the driver."], base.SCREENSHOTS_DIR / "screenshot-booking-details.png", None, "Figure 10 - Booking details")], "Tip - Check the detail page before departure to confirm the address, timing, and ride status."),
        base.Chapter("Features and functions - Offer and manage my trips", "Drivers can publish rides, define ride data, choose addresses, and manage passengers from their dedicated area.", [base.StepGroup("Open my trips", [f"Open {tr('shell.myTrips')} to view current, upcoming, and past rides.", f"Use {tr('driverTrips.publishNew')} to create a new ride."], base.SCREENSHOTS_DIR / "screenshot-driver-trips-list.png", None, "Figure 11 - Driver trips list"), base.StepGroup("Publish a new ride", ["Enter the date, time, seats, and departure and arrival addresses.", "Choose suggested addresses before you validate the publication."], base.SCREENSHOTS_DIR / "screenshot-driver-publish-trip.png", None, "Figure 12 - Publish trip form"), base.StepGroup("Monitor a published ride", ["Open an existing ride to review passengers and available actions.", "Use the detailed page to coordinate departure or cancel when necessary."], base.SCREENSHOTS_DIR / "screenshot-driver-trip-details.png", None, "Figure 13 - Driver trip management")], "Note - Publishing depends on valid address suggestions and the permissions available on the driver account."),
        base.Chapter("Features and functions - Messages and coordination", "The built-in chat helps drivers and passengers coordinate the practical details of a ride without leaving the app.", [base.StepGroup("Open the inbox", [f"Open {tr('shell.chats')} to display available conversations.", "Find new messages and open the conversation you need."], base.SCREENSHOTS_DIR / "screenshot-chat-inbox.png", None, "Figure 14 - Inbox"), base.StepGroup("Send and manage messages", [f"Type your message and use {tr('common.send')}.", "You can also select, copy, or clear messages from the visible conversation options."], base.SCREENSHOTS_DIR / "screenshot-chat-conversation.png", None, "Figure 15 - Open conversation")], "Tip - Use chat to confirm the exact pickup place and meeting time shortly before departure."),
        base.Chapter("Features and functions - My account and vehicle", "The personal account area lets users update profile data, strengthen login security, and maintain vehicle details used for driver rides.", [base.StepGroup("Update the profile", [f"Open {tr('shell.account')} and keep the {tr('common.profile')} tab selected.", "Edit visible personal information and save the changes."], base.SCREENSHOTS_DIR / "screenshot-account-profile.png", None, "Figure 16 - Profile tab"), base.StepGroup("Update the vehicle", [f"Switch to the {tr('common.car')} tab.", "Create or edit vehicle details and save them."], base.SCREENSHOTS_DIR / "screenshot-account-car.png", None, "Figure 17 - Car tab")], "Note - This area also contains sensitive actions such as password changes and account deletion."),
        base.Chapter("Troubleshooting", "This section helps users resolve common issues without contacting support. It focuses on the most frequent situations related to sign-in, ride search, and missing data.", [base.StepGroup("Check common issues", ["If a page does not load correctly, refresh the tab or sign in again.", "If a list is empty, verify that your account actually has a ride, booking, or conversation to display.", "If address suggestions do not appear, review your network connection and try the search again."], base.SCREENSHOTS_DIR / "screenshot-find-trip-results-empty.png", None, "Figure 18 - Example of a screen with limited or empty results")], "Tip - Most display issues come from missing data, an expired session, or unstable network access."),
        base.Chapter("FAQs", "This frequently asked questions section answers recurring user needs and reduces the need for direct support.", [base.StepGroup("Common questions", ["Why are some pages missing? Some areas depend on account permissions or available data.", "Why is a button disabled? One or more required fields may still be empty or invalid.", "Can I edit my personal information? Yes, from the My Account page."], base.SCREENSHOTS_DIR / "screenshot-account-profile.png", None, "Figure 19 - Account page used for common profile questions")], "Note - For more specific cases, use the Contact and support section at the end of this manual."),
        base.Chapter("Maintenance and care", "As a software application, maintenance mainly means keeping your information current, checking your password regularly, and maintaining clear ride and message data.", [base.StepGroup("Maintain the account", ["Update your profile if your displayed name, username, or phone number changes.", "Review vehicle details before publishing a new ride."], base.SCREENSHOTS_DIR / "screenshot-account-profile.png", None, "Figure 20 - Account maintenance area")], "Tip - A complete and current account helps other members recognize and trust you more easily."),
        base.Chapter("Contact and support information", "If you still need help, prepare a few key details before asking for support. This helps the response stay faster and more precise.", [base.StepGroup("Prepare a support request", ["Note the page, the time of the issue, and the action you were trying to complete.", "Keep a screenshot of the issue or visible message whenever possible.", "Share these details with your support contact or the organization that gave you access to the app."], base.SCREENSHOTS_DIR / "screenshot-home-dashboard.png", None, "Figure 21 - Example page to reference in a support request")], "Note - If your Covoit deployment does not display a direct support contact, reach out to your organization or project team."),
        base.Chapter("Appendix and glossary", "The appendix summarizes a few useful application terms so that all users can read the guide more easily.", [base.StepGroup("Quick glossary", ["Ride: a trip published by a driver.", "Booking: a passenger seat reserved on a ride.", "Conversation: a message thread between driver and passenger.", "My Account: the area used to manage personal and vehicle information."], base.SCREENSHOTS_DIR / "screenshot-home-dashboard.png", None, "Figure 22 - Main application areas used as glossary references")], "Tip - Revisit this glossary when interface terms feel unclear during early usage."),
        base.Chapter("Index", "This topic index helps readers find the most important manual subjects quickly, especially in a printed or exported format.", [base.StepGroup("Quick references", ["Account: see Introduction, Maintenance and care, Contact and support information.", "Sign-in: see Getting started.", "Messages: see Features and functions - Messages and coordination.", "Bookings: see Features and functions - My bookings.", "Driver trips: see Features and functions - Offer and manage my trips."], base.SCREENSHOTS_DIR / "screenshot-home-dashboard.png", None, "Figure 23 - Reference screen for the main manual topics")], "Note - The automatic Word table of contents remains the fastest way to jump to each section."),
    ]


def build_manual(locale: str, translations: dict[str, str], brand, version: str):
    tr = tr_factory(translations)
    document = Document()
    base.customize_styles(document, brand)
    base.build_cover(document, locale, translations["app.name"], version)
    base.add_toc(document, locale)

    chapters = build_manual_fr(tr) if locale == "fr" else build_manual_en(tr)
    embedded = 0
    placeholders = 0
    for chapter in chapters:
        e_count, p_count = base.add_chapter(document, chapter)
        embedded += e_count
        placeholders += p_count
        document.add_page_break()

    output = base.DOCS_DIR / ("guide-fr-manual.docx" if locale == "fr" else "guide-en-manual.docx")
    document.save(output)
    return {
        "locale": locale,
        "path": str(output),
        "chapters": len(chapters),
        "screenshots_embedded": embedded,
        "placeholders": placeholders,
        "validation": base.validate_guide(output),
    }


def main() -> None:
    package = base.load_json(base.PACKAGE_FILE)
    dictionaries = base.parse_dictionaries()
    brand = base.parse_brand_tokens()
    base.prepare_screenshots(brand)

    reports = [
        build_manual("fr", dictionaries["fr"], brand, package["version"]),
        build_manual("en", dictionaries["en"], brand, package["version"]),
    ]

    summary = {
        "reports": reports,
        "manual_review": [
            "Refresh the Word table of contents to update final page numbers.",
            "Review wording if your deployment uses organization-specific legal or support policies.",
        ],
    }
    print(json.dumps(summary, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
